"""Local, deterministic text extraction for the learning workbench.

This module reports literal matches only. It does not assign a match to a
business field or decide whether a document is correct.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path
import re
from typing import Iterable


@dataclass(frozen=True)
class Candidate:
    kind: str
    value: str
    location: str
    context: str

    def to_dict(self) -> dict[str, str]:
        return asdict(self)


@dataclass(frozen=True)
class ExtractedDocument:
    path: str
    pages: int
    text: str
    candidates: tuple[Candidate, ...]
    warning: str = ""


PATTERNS: tuple[tuple[str, re.Pattern[str]], ...] = (
    ("Email", re.compile(r"(?<![\w.+-])[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}(?![\w.-])")),
    ("Protocollo", re.compile(r"\b(?:prot(?:ocollo)?\.?\s*(?:n\.?|numero)?\s*[:#-]?\s*)[A-Z0-9][A-Z0-9/.-]{2,}\b", re.I)),
    ("Importo", re.compile(r"(?<!\w)(?:€\s*)?\d{1,3}(?:[. ]\d{3})*(?:,\d{1,2})?\s*(?:€|euro)?(?!\w)", re.I)),
    ("Data", re.compile(r"\b(?:\d{1,2}[./-]\d{1,2}[./-](?:\d{2}|\d{4})|\d{1,2}\s+(?:gennaio|febbraio|marzo|aprile|maggio|giugno|luglio|agosto|settembre|ottobre|novembre|dicembre)\s+\d{4})\b", re.I)),
    ("Orario", re.compile(r"\b(?:[01]?\d|2[0-3])[:.]\d{2}\s*(?:-|–|—|alle)\s*(?:[01]?\d|2[0-3])[:.]\d{2}\b", re.I)),
)


def _context(text: str, start: int, end: int, radius: int = 65) -> str:
    lo, hi = max(0, start - radius), min(len(text), end + radius)
    snippet = " ".join(text[lo:hi].split())
    return ("…" if lo else "") + snippet + ("…" if hi < len(text) else "")


def find_candidates(text: str, locations: Iterable[tuple[str, str]]) -> tuple[Candidate, ...]:
    """Find exact regex occurrences; locations is (text, page/section label)."""
    found: list[Candidate] = []
    for section_text, location in locations:
        for kind, pattern in PATTERNS:
            for match in pattern.finditer(section_text):
                found.append(Candidate(kind, match.group(0).strip(), location, _context(section_text, match.start(), match.end())))
    return tuple(found)


def read_document(path: str | Path) -> ExtractedDocument:
    """Read supported text-bearing files locally, retaining page provenance."""
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    sections: list[tuple[str, str]] = []
    warning = ""

    if suffix == ".pdf":
        try:
            import fitz  # PyMuPDF
        except ImportError as exc:  # pragma: no cover - exercised in packaged app
            raise RuntimeError("Lettura PDF non disponibile: PyMuPDF non è installato.") from exc
        with fitz.open(file_path) as pdf:
            for index, page in enumerate(pdf, start=1):
                sections.append((page.get_text("text"), f"pagina {index}"))
            page_count = len(pdf)
        if not any(text.strip() for text, _ in sections):
            warning = "Il PDF non contiene testo selezionabile. OCR non incluso in questa versione; trascrivere o annotare manualmente."
    elif suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("Lettura DOCX non disponibile: python-docx non è installato.") from exc
        doc = Document(file_path)
        lines: list[str] = []
        lines.extend(p.text for p in doc.paragraphs if p.text.strip())
        for table_index, table in enumerate(doc.tables, start=1):
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        sections.append(("\n".join(lines), "documento DOCX"))
        page_count = 1
    elif suffix in {".txt", ".md", ".csv"}:
        content = file_path.read_text(encoding="utf-8-sig", errors="replace")
        sections.extend((chunk, f"blocco {i + 1}") for i, chunk in enumerate(content.split("\f")))
        page_count = len(sections)
    else:
        raise ValueError(f"Formato non supportato: {suffix or '(senza estensione)'}. Formati supportati: PDF con testo, DOCX, TXT, MD, CSV.")

    text = "\n\n".join(part for part, _ in sections)
    candidates = find_candidates(text, sections)
    return ExtractedDocument(str(file_path), page_count, text, candidates, warning)

